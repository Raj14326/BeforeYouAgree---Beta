from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOWNLOADS = Path(r"C:\Users\Raj\Downloads")
PAIR_SOURCE = DOWNLOADS / "Pair_Programming_Observations_Iteration_1.docx"
PAIR_OUTPUT = DOWNLOADS / "Pair_Programming_Observations_Iteration_2.docx"
TEST_SOURCE = DOWNLOADS / "Interation1_Testing.docx"
TEST_OUTPUT = DOWNLOADS / "Iteration2_Testing.docx"


def set_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def set_cell(cell, value):
    # cell.text intentionally removes outdated embedded commit screenshots.
    cell.text = value


def add_table(document, rows, style="Table Grid"):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        table.style = style
    except KeyError:
        properties = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "B7B7B7")
            borders.append(border)
        properties.append(borders)
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    return table


# ---------------------------------------------------------------------------
# Pair-programming / collaboration record
# ---------------------------------------------------------------------------
pair = Document(PAIR_SOURCE)
set_text(pair.paragraphs[1], (
    "Iteration 2\nThis record uses Git history from Prototype and feat/will. Git confirms "
    "authors, changes and merges, but it does not prove a live pairing session. Confirm the "
    "navigator and duration from team notes before submission."
))

summary = [
    ["Area", "Main contributor", "Git-supported collaborator", "Primary commit"],
    ["LegalBERT categories", "Will", "Leo model handover; SH4LANTW merge review", "a680188e"],
    ["Clause severity levels", "Will", "SH4LANTW merged the UI refinement", "0387a1dc"],
    ["Preference-based UI", "Will", "SH4LANTW merged the feature work", "d978ccfa"],
    ["Category filtering", "Will", "SH4LANTW merged PR #12", "a982c976"],
    ["Personalised scoring", "Raj", "Built on Will's preference UI; later refined on feat/will", "6dac8f62"],
]
for row, values in zip(pair.tables[0].rows, summary):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value)

sessions = [
    {
        "heading": "LegalBERT model and risk categories",
        "date": "12 September 2026",
        "driver": "Will (commit author)",
        "collaborator": "Leo supplied the trained model; SH4LANTW reviewed/merged related model work",
        "task": "Replace M006 with the LegalBERT multi-label model and expose category labels in the UI.",
        "completed": "Converted the model to ONNX, added server/bert-model.ts, changed /api/analyze to LegalBERT and displayed detected categories.",
        "decision": "The binary M006 result could not support category-based preferences. The team moved to an eight-category LegalBERT response.",
        "learned": "The backend response, thresholds and frontend category types must change together.",
        "evidence": "a680188e on Prototype and feat/will",
    },
    {
        "heading": "Clear clause severity levels",
        "date": "15 September 2026",
        "driver": "Will (commit author)",
        "collaborator": "SH4LANTW merged the UI refinement into Prototype",
        "task": "Replace the binary risk badge with low, medium and high severity levels.",
        "completed": "Added server/risk-level.ts, severity messages, review categories, UI badges and risk-level unit tests.",
        "decision": "Severity is based on category thresholds and detection strength, not a raw score presented as legal probability.",
        "learned": "Severity rules need boundary tests so low, medium and high remain consistent.",
        "evidence": "0387a1dc on feat/will; merged through e66445ec",
    },
    {
        "heading": "Five-region interface and preference sidebar",
        "date": "15 September 2026",
        "driver": "Will (commit author)",
        "collaborator": "SH4LANTW merged the feature branch",
        "task": "Separate the interface into search, service documents, clauses, original text and risk preferences.",
        "completed": "Split App.vue into reusable components and added the preference sidebar and supporting frontend libraries.",
        "decision": "The previous single component was difficult to change and test. Smaller components gave each feature a clear responsibility.",
        "learned": "Component boundaries make preference and clause behaviour easier to test independently.",
        "evidence": "d978ccfa on feat/will; merged through e66445ec",
    },
    {
        "heading": "Risk-category toggles and clause filtering",
        "date": "16 September 2026",
        "driver": "Will (commit author)",
        "collaborator": "SH4LANTW merged PR #12",
        "task": "Allow users to enable or disable risk categories and filter clause cards.",
        "completed": "Added category definitions, preference switches, privacy grouping and enabled-category filtering in ClausesPanel.",
        "decision": "A clause remains visible when at least one of its detected categories is enabled.",
        "learned": "Filtering rules must handle clauses with several categories without hiding valid matches.",
        "evidence": "a982c976 on feat/will; merged through 5777de86",
    },
    {
        "heading": "Preference ordering and personalised risk score",
        "date": "16 September 2026",
        "driver": "Raj (commit author)",
        "collaborator": "Built on Will's preference components; Will later refined the preference UI",
        "task": "Reorder findings using user priorities and show a personalised risk score.",
        "completed": "Added draggable category priority, priority-based clause ordering, personalised score calculation and component/unit tests.",
        "decision": "Detection strength remains the main score, while preference order adds a smaller relevance bonus.",
        "learned": "A personalised score is a relevance score, not the probability of legal harm, so the UI must label it clearly.",
        "evidence": "6dac8f62 on Prototype and feat/will",
    },
]

starts = [4, 17, 28, 40, 51]
for start, session in zip(starts, sessions):
    values = [
        session["heading"],
        f"Date: {session['date']}",
        "Duration: Confirm from team session notes",
        f"Driver / main contributor: {session['driver']}",
        f"Navigator / collaborator: {session['collaborator']}",
        f"Task: {session['task']}",
        f"Completed: {session['completed']}",
        f"Problem and decision: {session['decision']}",
        f"What was learned: {session['learned']}",
        f"Git evidence: {session['evidence']}",
    ]
    for offset, value in enumerate(values):
        set_text(pair.paragraphs[start + offset], value)

for index, table in enumerate(pair.tables[1:], start=1):
    set_cell(table.rows[0].cells[0], f"Session evidence: {sessions[index - 1]['evidence']}")

# Remove the old Iteration 1 screenshots. Their commits no longer match the
# updated Iteration 2 sessions; the verified hashes remain in the evidence boxes.
for paragraph in pair.paragraphs:
    for drawing in list(paragraph._p.iter(qn("w:drawing"))):
        drawing.getparent().remove(drawing)
    for picture in list(paragraph._p.iter(qn("w:pict"))):
        picture.getparent().remove(picture)
for table in pair.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for drawing in list(paragraph._p.iter(qn("w:drawing"))):
                    drawing.getparent().remove(drawing)
                for picture in list(paragraph._p.iter(qn("w:pict"))):
                    picture.getparent().remove(picture)

pair.core_properties.title = "Before You Agree - Pair Programming and Collaboration Record - Iteration 2"
pair.save(PAIR_OUTPUT)


# ---------------------------------------------------------------------------
# Iteration 2 TDD and acceptance-testing document
# ---------------------------------------------------------------------------
testing = Document(TEST_SOURCE)
body = testing._body._element
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

testing.add_paragraph("Iteration 2 Testing and TDD Evidence", style="Heading 1")
testing.add_paragraph(
    "Scope: risk preferences, category filtering, finding order, category labels, fallback behaviour and severity levels.",
    style="Normal",
)
testing.add_paragraph("Testing approach", style="Heading 2")
testing.add_paragraph(
    "Iteration 2 uses a TDD cycle: write the expected behaviour, run the test to see it fail, add the smallest change, then refactor while keeping the suite green. Vitest and Vue Test Utils cover component and scoring logic. Playwright is used for browser-session behaviour.",
)
testing.add_paragraph("Current automated result", style="Heading 2")
testing.add_paragraph("Command: npx vitest run RiskPreferenceSidebar, ClausesPanel, ClauseCard, personalised-risk-score and risk-level specs", style="HTML Preformatted")
testing.add_paragraph("Result on 17 September 2026: 5 test files passed; 19 tests passed.", style="HTML Preformatted")
testing.add_paragraph(
    "Important: the full suite still has four failing legacy LegalBERT adapter tests. Criteria marked Gap or Red below must not be reported as passed until their tests and implementation are completed.",
)

testing.add_paragraph("Acceptance-criteria traceability", style="Heading 2")
add_table(testing, [
    ["AC", "Behaviour", "Test evidence", "Status"],
    ["4.2.1", "Reorder findings without reload", "ClausesPanel.spec.ts", "Pass"],
    ["4.2.2", "Prioritise selected categories", "ClausesPanel and personalised score specs", "Pass"],
    ["4.1.1", "Display preference toggles", "RiskPreferenceSidebar.spec.ts covers master control", "Partial"],
    ["4.1.2", "Highlight/filter by selected category", "Dedicated acceptance test still required", "Gap"],
    ["4.3.1", "Reset after browser session", "Playwright session test required", "Gap"],
    ["4.3.2", "Keep preferences in same session", "Playwright navigation test required", "Gap"],
    ["5.2.2", "Clear visual severity difference", "ClauseCard and score specs", "Pass"],
    ["5.1.1", "Assign category to each flagged clause", "LegalBERT/category code; full adapter suite not green", "Partial"],
    ["5.1.2", "Fallback category for unclear result", "Failing test must be written; fallback not present", "Red"],
    ["5.2.1", "Assign low/medium/high severity", "server/risk-level.spec.ts", "Pass"],
])

cases = [
    ("AC 4.2.1 - Reorder findings without reloading", "Given two findings with different category priorities; when preferences are enabled and categoryPriority changes; then the clause-card order changes immediately and no API reload is made.", "src/components/ClausesPanel.spec.ts", "Pass", "expect(clauseTexts()).toEqual(['Priority category but lower score', 'Lower priority category but higher score'])"),
    ("AC 4.2.2 - Prioritise selected risk categories by severity", "Given a strong lower-priority match and a weaker first-priority match; when preferences are on; then the selected category appears first. When preferences are off, findings use descending detection score.", "ClausesPanel.spec.ts and personalised-risk-score.spec.ts", "Pass", "expect(scoreForFirstPreference).toBeGreaterThan(scoreForLastPreference)"),
    ("AC 4.1.1 - Display risk preference toggles", "Given the preference sidebar; when it opens; then every configured category has a labelled switch and the master preference control is visible.", "RiskPreferenceSidebar.spec.ts", "Partial: master behaviour is tested; add an assertion for all category switches", "expect(wrapper.findAll('input[role=switch]').length).toBeGreaterThan(1)"),
    ("AC 4.1.2 - Highlight clauses using the selected category", "Given clauses in two categories; when one category is disabled; then clauses that only match it are hidden and the enabled category remains highlighted in the source document.", "Add ClausesPanel/Vue integration test", "Gap", "disable('arbitration'); expect(cards()).not.toContain(arbitrationClause)"),
    ("AC 4.3.1 - Reset preferences after the browser session", "Given changed preferences; when a new browser context is created; then the default enabled categories and order are restored.", "Add Playwright test with two browser contexts", "Gap", "close context; open new context; expect(defaultPreferences).toBeVisible()"),
    ("AC 4.3.2 - Keep preferences while browsing in the same session", "Given changed toggles and order; when the user selects another service without reloading the application; then those preferences remain unchanged.", "Add Playwright service-navigation test", "Gap", "select service B; expect(savedToggle).toBeChecked(); expect(order).toEqual(chosenOrder)"),
    ("AC 5.2.2 - Clear visual distinction by severity", "Given low, medium and high scores; when clause cards render; then each level has the correct label, badge class and colour treatment.", "ClauseCard.spec.ts and personalised-risk-score.spec.ts", "Pass", "expect(personalisedRiskLevel(44, 45, 75)).to map to low, medium and high boundaries"),
    ("AC 5.1.1 - Assign a risk category to each flagged clause", "Given a clause above a LegalBERT category threshold; when analysis runs; then the finding contains at least one category ID and name and the card displays it.", "server/bert-model.spec.ts plus ClauseCard test", "Partial: category behaviour exists, but four adapter tests in the full suite fail", "expect(finding.categories.length).toBeGreaterThan(0)"),
    ("AC 5.1.2 - Use a fallback category for unclear classification", "Given a clause marked risky without a usable category; when the result is normalised; then it receives a named fallback category such as General risk.", "New backend normalisation test required", "Red: fallback behaviour is not currently implemented", "expect(finding.categories).toContainEqual({ id: 'general_risk', name: 'General risk' })"),
    ("AC 5.2.1 - Assign a clear severity level", "Given scores below, near and clearly above category thresholds; when classifyRiskLevel runs; then it returns low, medium or high using the documented boundaries.", "server/risk-level.spec.ts", "Pass", "expect low for no match, medium for one/near match, high for strong or multiple matches"),
]

for heading, behaviour, evidence, status, code in cases:
    testing.add_paragraph(heading, style="Heading 2")
    testing.add_paragraph(f"Test case: {behaviour}")
    testing.add_paragraph(f"Evidence: {evidence}")
    testing.add_paragraph(f"Status: {status}")
    testing.add_paragraph(code, style="HTML Preformatted")

testing.add_paragraph("TDD follow-up", style="Heading 2")
testing.add_paragraph(
    "The next red tests should cover AC 4.1.2, AC 4.3.1, AC 4.3.2 and AC 5.1.2. The fallback-category test should be written before adding the fallback implementation. After these pass, rerun the full unit and Playwright suites and attach the output to this document.",
)

testing.core_properties.title = "Before You Agree - Iteration 2 Testing and TDD Evidence"
testing.save(TEST_OUTPUT)

print(PAIR_OUTPUT)
print(TEST_OUTPUT)
