from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOWNLOADS = Path(r"C:\Users\Raj\Downloads")
SOURCE = DOWNLOADS / "Interation1_Testing.docx"
OUTPUT = DOWNLOADS / "Iteration2_Playwright_Testing.docx"


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "B7B7B7")
        borders.append(border)
    properties.append(borders)
    for row_index, (row, values) in enumerate(zip(table.rows, rows)):
        for cell, value in zip(row.cells, values):
            run = cell.paragraphs[0].add_run(value)
            if row_index == 0:
                run.bold = True
    return table


def add_case(document, heading, purpose, steps, expected, code):
    document.add_paragraph(heading, style="Heading 2")
    document.add_paragraph(f"Purpose: {purpose}")
    document.add_paragraph(f"Steps: {steps}")
    document.add_paragraph(f"Expected result: {expected}")
    for line in code.strip().splitlines():
        document.add_paragraph(line, style="HTML Preformatted")


doc = Document(SOURCE)
body = doc._body._element
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

doc.add_paragraph("Iteration 2 Playwright Acceptance Testing", style="Heading 1")
doc.add_paragraph("Acceptance criteria 4.1.1–4.3.2 and 5.1.1–5.2.2", style="Heading 1")

doc.add_paragraph("Testing approach", style="Heading 2")
doc.add_paragraph(
    "The Iteration 2 acceptance tests use Playwright with controlled API responses. Each test opens the real Vue interface in Chromium and checks the same actions a user performs. The API is mocked so service data, clause categories and severity levels remain repeatable."
)
doc.add_paragraph(
    "The tests follow a TDD cycle. The expected browser behaviour was written first. The first run exposed missing category-based document highlights and a missing fallback label. Small UI changes were then added until the acceptance suite passed."
)

doc.add_paragraph("Commands and results", style="Heading 2")
doc.add_paragraph(
    "npx playwright test e2e/iteration-2-acceptance.spec.ts --project=chromium",
    style="HTML Preformatted",
)
doc.add_paragraph("Running 8 tests using 1 worker", style="HTML Preformatted")
doc.add_paragraph("8 passed (42.4s)", style="HTML Preformatted")
doc.add_paragraph("npm run build", style="HTML Preformatted")
doc.add_paragraph("Type-check and production build passed.", style="HTML Preformatted")

doc.add_paragraph("TDD evidence", style="Heading 2")
add_table(doc, [
    ["Stage", "Evidence"],
    ["Red", "The first Playwright run reported 5 passed and 3 failed. It exposed the missing fallback category and selected-category highlight behaviour; two test selectors were also corrected."],
    ["Green", "A General risk fallback and category-aware document highlights were added. The second run passed all 8 tests."],
    ["Refactor", "The highlight rule was added as an optional document-view predicate, keeping existing callers unchanged. The production build then passed."],
])

doc.add_paragraph("Acceptance-criteria traceability", style="Heading 2")
add_table(doc, [
    ["AC", "Playwright test", "Result"],
    ["4.1.1", "Displays labelled risk preference toggles", "Pass"],
    ["4.1.2", "Filters clause cards and document highlights", "Pass"],
    ["4.2.1", "Reorders findings without reload or re-analysis", "Pass"],
    ["4.2.2", "Uses preference priority and severity fallback order", "Pass"],
    ["4.3.1", "Resets preferences in a new browser context", "Pass"],
    ["4.3.2", "Keeps preferences while changing services", "Pass"],
    ["5.1.1", "Shows the model category on each categorised clause", "Pass"],
    ["5.1.2", "Shows General risk when a risky clause has no category", "Pass"],
    ["5.2.1", "Shows low, medium and high severity labels", "Pass"],
    ["5.2.2", "Uses different severity strip colours", "Pass"],
])

add_case(
    doc,
    "AC 4.1.1 - Display risk preference toggles",
    "Confirm that the sidebar provides accessible preference controls.",
    "Open the application and inspect the master switch and category switches.",
    "The master, Unilateral change and Arbitration switches are visible and checked, with at least ten labelled switches available.",
    """
await expect(page.getByLabel('Enable risk preferences')).toBeChecked()
await expect(page.getByLabel('Include Unilateral change in risk preferences')).toBeChecked()
await expect(page.getByLabel('Include Arbitration in risk preferences')).toBeChecked()
expect(await page.getByRole('switch').count()).toBeGreaterThanOrEqual(10)
""",
)

add_case(
    doc,
    "AC 4.1.2 - Highlight clauses using the selected category",
    "Check that disabling a category updates both the finding list and original document.",
    "Analyse a document, open the full text, and disable Arbitration.",
    "The arbitration card and highlight disappear while other findings remain.",
    """
const documentMarks = page.locator('#original-document-view mark.clause-mark')
await expect(documentMarks).toHaveCount(3)
await page.getByLabel('Include Arbitration in risk preferences').uncheck()
await expect(page.locator('.clause-card-text')).not.toContainText([arbitrationClause])
await expect(documentMarks).toHaveCount(2)
""",
)

add_case(
    doc,
    "AC 4.2.1 - Reorder findings without reloading the page",
    "Confirm that drag-and-drop changes finding order without another analysis request.",
    "Analyse once, drag Arbitration above Unilateral change, and inspect the first clause.",
    "Arbitration becomes the first finding and the analysis request count remains one.",
    """
await arbitration.locator('.preference-drag-handle')
  .dragTo(unilateral.locator('.preference-drag-handle'))
await expect(page.locator('.clause-card-text').first()).toHaveText(arbitrationClause)
expect(counters.analyze).toBe(1)
""",
)

add_case(
    doc,
    "AC 4.2.2 - Prioritise selected categories by severity",
    "Check preference order when enabled and objective score order when disabled.",
    "Analyse the fixture, check the first finding, then turn preferences off.",
    "The preferred category appears first when enabled; the stronger detection appears first when disabled.",
    """
await expect(page.locator('.clause-card-text').first()).toHaveText(unilateralClause)
await page.getByLabel('Enable risk preferences').uncheck()
await expect(page.locator('.clause-card-text').first()).toHaveText(arbitrationClause)
""",
)

add_case(
    doc,
    "AC 4.3.1 - Reset preferences after the browser session",
    "Confirm that preferences are not stored permanently.",
    "Disable Arbitration in one browser context, close it, and open a new context.",
    "Arbitration returns to its default checked state in the new session.",
    """
await firstPage.getByLabel('Include Arbitration in risk preferences').uncheck()
await first.close()
const second = await browser.newContext()
await expect(secondPage.getByLabel('Include Arbitration in risk preferences')).toBeChecked()
""",
)

add_case(
    doc,
    "AC 4.3.2 - Keep preferences while browsing services",
    "Confirm that preference state remains while the same page session is active.",
    "Disable Arbitration, select Google, and inspect the same switch.",
    "The selected service changes but Arbitration remains disabled.",
    """
await page.getByLabel('Include Arbitration in risk preferences').uncheck()
await page.getByLabel('Service').fill('Goo')
await page.locator('.list-group-item').filter({ hasText: 'Google' }).click()
await expect(page.getByLabel('Include Arbitration in risk preferences')).not.toBeChecked()
""",
)

add_case(
    doc,
    "AC 5.1.1 and 5.1.2 - Category and fallback labels",
    "Confirm that model categories are shown and unclear risky results have a fallback.",
    "Analyse clauses with Unilateral change, Arbitration and no supplied category.",
    "The first two show their model category; the uncategorised risky clause shows General risk.",
    """
await expect(card(unilateralClause)).toContainText('Unilateral change')
await expect(card(arbitrationClause)).toContainText('Arbitration')
await expect(card(unclearClause)).toContainText('General risk')
""",
)

add_case(
    doc,
    "AC 5.2.1 and 5.2.2 - Clear severity levels",
    "Confirm that each clause has a severity label and an objective visual difference.",
    "Render high, medium and low fixture findings and read their labels and strip colours.",
    "All three labels appear and the three clause strips use different computed colours.",
    """
await expect(card(unilateralClause)).toContainText('Medium personalised risk')
await expect(card(arbitrationClause)).toContainText('High personalised risk')
await expect(card(unclearClause)).toContainText('Low personalised risk')
expect(new Set(stripColors).size).toBe(3)
""",
)

doc.add_paragraph("Scope and limitations", style="Heading 2")
doc.add_paragraph(
    "The suite uses controlled API fixtures and therefore tests frontend acceptance behaviour, not live ToS;DR, GitHub, S3 or LegalBERT availability. The tests run in Chromium. The existing full unit suite still contains four legacy LegalBERT adapter failures that are tracked separately and are not hidden by this acceptance result."
)

doc.core_properties.title = "Before You Agree - Iteration 2 Playwright Acceptance Testing"
doc.core_properties.comments = "Generated from the passing e2e/iteration-2-acceptance.spec.ts suite."
doc.save(OUTPUT)
print(OUTPUT)
