from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\Raj\Downloads\Iteration 1 build document.md.docx")
OUTPUT = Path(r"C:\Users\Raj\Downloads\Iteration 2 build document.docx")


def set_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def set_cell(cell, value):
    set_text(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        set_text(paragraph, "")


def fill_table(table, rows):
    while len(table.rows) < len(rows):
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)


doc = Document(SOURCE)

# Technology stack: update only the backend and analysis-model rows.
stack = doc.tables[1]
set_cell(stack.rows[5].cells[1], "Node.js API + Nginx")
set_cell(
    stack.rows[5].cells[2],
    "Node.js runs the API. Nginx receives web traffic and forwards API requests to Node.js.",
)
set_cell(stack.rows[6].cells[1], "LegalBERT Small ONNX classifier")
set_cell(
    stack.rows[6].cells[2],
    "A local multi-label model that checks each clause against eight unfair-term categories, supported by privacy rules.",
)

# Major files and folders: replace the old M006-era list with the current project structure.
fill_table(doc.tables[2], [
    ["Path", "What it holds"],
    ["index.html", "The browser page shell, early theme setup and frontend asset links."],
    ["src/main.ts", "Loads the main styles and starts the Vue application."],
    ["src/App.vue", "Main application layout and coordination between search, documents, analysis and preferences."],
    ["src/types.ts", "Shared frontend types for services, documents, clauses, categories and API results."],
    ["src/components/SearchBar.vue", "Service search, keyboard controls and short-term search-result caching."],
    ["src/components/ServiceDocumentCard.vue", "Displays the selected service and its available policy documents."],
    ["src/components/ClausesPanel.vue", "Displays analysed clauses and applies category filters."],
    ["src/components/ClauseCard.vue", "Shows one clause, its categories, risk level and personalised score."],
    ["src/components/OriginalDocumentPanel.vue", "Shows the complete source document with highlighted findings."],
    ["src/components/RiskPreferenceSidebar.vue", "Lets the user enable, disable and order risk-category preferences."],
    ["src/components/QuickGuide.vue", "Provides the in-application usage guide."],
    ["src/assets/main.css", "Custom styling, responsive layout, themes and clause highlighting."],
    ["src/assets/BYA_logo.png", "The Before You Agree logo."],
    ["src/lib/api.ts", "Builds API addresses and contains shared frontend API helpers."],
    ["src/lib/document-view.ts", "Prepares document text and safe highlight markup for display."],
    ["src/lib/risk-categories.ts", "Frontend definitions and descriptions for risk categories."],
    ["src/lib/risk-level.ts", "Frontend risk-level display helpers."],
    ["src/lib/personalised-risk-score.ts", "Combines detection strength with the user's category preferences."],
    ["server/index.ts", "Main API server, route handling, CORS, input limits, caching, rate limiting and external requests."],
    ["server/bert-model.ts", "Loads the local LegalBERT ONNX model and produces clause-level category findings."],
    ["server/clauses.ts", "Splits documents into clauses and keeps their exact text positions."],
    ["server/privacy-rules.ts", "Adds clear rule-based findings for important privacy-policy risks."],
    ["server/risk-level.ts", "Assigns low, medium or high levels from model scores and thresholds."],
    ["server/model-storage.ts", "Downloads missing private model files from S3 during deployment."],
    ["server/html-to-plain-text.ts", "Converts policy HTML into clean plain text."],
    ["server/*.spec.ts and src/**/*.spec.ts", "Unit tests for the API, model adapter, privacy rules, components and risk scoring."],
    ["ml/local-models/bya-legalbert-small-unfair-tos/", "Local tokenizer, risk configuration and ONNX model files used for inference."],
    ["e2e/", "Playwright acceptance tests for the main user stories and quick guide."],
    ["deploy/ec2-post-deploy.sh", "Restarts the EC2 API after deployment and checks /api/health."],
    ["Dockerfile.aws", "Production container build for the Node.js API and local model."],
    ["amplify.yml", "Frontend build and AWS Amplify hosting settings."],
    ["apprunner.yaml", "Alternative AWS App Runner build and runtime settings."],
    ["vite.config.ts", "Frontend build settings and the development /api proxy."],
    ["vitest.config.ts", "Unit-test configuration."],
    ["playwright.config.ts", "Local browser-test configuration."],
    ["playwright.deployed.config.ts", "Browser-test configuration for the deployed application."],
    ["package.json and package-lock.json", "Dependencies, scripts and locked package versions."],
    ["README.md", "Technical setup and project notes."],
])

# Data-source/model description.
set_text(
    doc.paragraphs[71],
    "The LegalBERT Small ONNX model (stored locally or downloaded from private S3) performs the risk analysis. It checks clauses against eight unfair-term categories, while rule-based privacy checks add clear privacy findings. No external AI inference service is required.",
)

# Current risk-analysis process.
set_text(
    doc.paragraphs[77],
    "The document is split into sentences and list items. Very short fragments are skipped, and repeated clauses are grouped while their original positions are kept.",
)
set_text(
    doc.paragraphs[78],
    "The LegalBERT tokenizer converts each clause into tokens. The local ONNX model processes the clauses in small batches and scores eight unfair-term categories.",
)
set_text(
    doc.paragraphs[79],
    "Each category score is compared with its own tuned threshold. Rule-based privacy checks are added, and a clause is marked risky when at least one category is detected.",
)
set_text(
    doc.paragraphs[80],
    "The application assigns a low, medium or high risk level. The frontend shows the categories, applies the user's preferences and highlights each finding in the full document.",
)
set_text(
    doc.paragraphs[82],
    "The results are automated predictions based on learned patterns and privacy rules. They may contain errors and must not be treated as legal advice.",
)

# Current backend deployment.
set_text(
    doc.paragraphs[88],
    "Backend API → deployed on AWS EC2 as a Node.js service behind Nginx. AWS CodePipeline and Systems Manager support deployment, and the post-deploy script checks GET /api/health.",
)

doc.core_properties.title = "Before You Agree - Iteration 2 Build Document"
doc.core_properties.comments = "Only the requested technology, files, model, analysis and backend deployment sections were updated."
doc.save(OUTPUT)
print(OUTPUT)
